from __future__ import annotations
import argparse, csv, json, re, time
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from lpi_core import determine_pan_ids,load_lpi_expression,fit_matrix,bh_adjust

SEED=20260915

def prism_family(target,moa):
    t='' if pd.isna(target) else str(target).upper(); m='' if pd.isna(moa) else str(moa).upper(); fam=[]
    if ('IGF-1 INHIBITOR' in m or 'INSULIN GROWTH FACTOR RECEPTOR INHIBITOR' in m or
        re.search(r'(^|[, ]+)IGF1R($|[, ]+)',t)): fam.append('IGF1R')
    if 'PI3K INHIBITOR' in m: fam.append('PI3K')
    if 'AKT INHIBITOR' in m: fam.append('AKT')
    if 'MTOR INHIBITOR' in m: fam.append('mTOR')
    return '|'.join(fam) if fam else np.nan

def gdsc_family(target):
    t='' if pd.isna(target) else str(target).upper(); fam=[]
    if 'IGF1R' in t: fam.append('IGF1R')
    if 'PI3K' in t or 'PIK3' in t: fam.append('PI3K')
    if re.search(r'(^|[^A-Z0-9])AKT([123]?)([^A-Z0-9]|$)',t): fam.append('AKT')
    if 'MTOR' in t: fam.append('mTOR')
    return '|'.join(fam) if fam else np.nan

def family_bh(df,eligible_col='eligible'):
    df=df.copy(); df['q_global']=np.nan; df['q_family']=np.nan
    ok=df[eligible_col].astype(bool) & np.isfinite(df.p)
    if ok.any(): df.loc[ok,'q_global']=bh_adjust(df.loc[ok,'p'].to_numpy())
    fam=ok & df.is_family.astype(bool)
    if fam.any(): df.loc[fam,'q_family']=bh_adjust(df.loc[fam,'p'].to_numpy())
    return df

def read_gdsc_xlsx(path: Path):
    wb=load_workbook(path,read_only=True,data_only=True); ws=wb.active; it=ws.iter_rows(values_only=True); header=next(it)
    col={str(x):i for i,x in enumerate(header)}
    need=['SANGER_MODEL_ID','DRUG_ID','DRUG_NAME','PUTATIVE_TARGET','PATHWAY_NAME','LN_IC50','AUC']
    missing=[x for x in need if x not in col]
    if missing: raise RuntimeError(f'{path.name} missing columns {missing}')
    rows=[]
    for r in it:
        rows.append(tuple(r[col[x]] for x in need))
    return pd.DataFrame(rows,columns=need)

def model_table(arr,labels,metadata,min_n,min_lineages):
    df=pd.DataFrame({'drug_id':labels,'beta':arr[:,0],'se':arr[:,1],'t':arr[:,2],'p':arr[:,3],
                     'n':arr[:,4].astype(int),'partial_r2':arr[:,5]})
    df=df.merge(metadata,on='drug_id',how='left')
    df['n_lineages']=df['n_lineages'].astype(int)
    df['family']=df['target'].map(gdsc_family); df['is_family']=df.family.notna()
    df['eligible']=(df.n>=min_n)&(df.n_lineages>=min_lineages)
    df=family_bh(df)
    return df

def normalized_drug(s): return re.sub(r'[^a-z0-9]','',str(s).lower())

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--data-dir',type=Path,required=True); ap.add_argument('--out-dir',type=Path,required=True)
    ap.add_argument('--min-n',type=int,default=100); ap.add_argument('--min-lineages',type=int,default=5); args=ap.parse_args(); args.out_dir.mkdir(parents=True,exist_ok=True)
    meta,header,expr_header,ids,crispr_path=determine_pan_ids(args.data_dir)
    _,lineages,lpi,pro=load_lpi_expression(args.data_dir,ids,meta)
    score=pd.DataFrame({'ModelID':ids,'LPI':lpi,'proliferation':pro,'lineage':lineages,
                        'SangerModelID':[meta[x].get('SangerModelID','') for x in ids]})
    score.to_csv(args.out_dir/'cell_line_scores.csv',index=False)

    pm=pd.read_csv(args.data_dir/'Repurposing_Public_24Q2_Extended_Primary_Data_Matrix.csv',index_col=0)
    ann=pd.read_csv(args.data_dir/'Repurposing_Public_24Q2_Extended_Primary_Compound_List.csv')
    ann=ann.drop_duplicates('IDs',keep='first').set_index('IDs')
    prism_ids=[x for x in ids if x in pm.columns]
    ix=np.array([ids.index(x) for x in prism_ids])
    Y=pm.loc[:,prism_ids].T.to_numpy(dtype=float)
    arr=fit_matrix(Y,lpi[ix],pro[ix],lineages[ix])
    pr=pd.DataFrame({'id':pm.index.astype(str),'n':arr[:,4].astype(int),'beta':arr[:,0],'se':arr[:,1],'t':arr[:,2],'p':arr[:,3],'partial_r2':arr[:,5]})
    pr['n_lineages']=[len(set(np.asarray(lineages)[ix][np.isfinite(Y[:,j])])) for j in range(Y.shape[1])]
    pr['screen']=[ann.loc[x,'screen'] if x in ann.index else np.nan for x in pr.id]
    pr['dose']=[ann.loc[x,'dose'] if x in ann.index else np.nan for x in pr.id]
    pr['target']=[ann.loc[x,'repurposing_target'] if x in ann.index else np.nan for x in pr.id]
    pr['moa']=[ann.loc[x,'MOA'] if x in ann.index else np.nan for x in pr.id]
    pr['drug']=[ann.loc[x,'Drug.Name'] if x in ann.index else np.nan for x in pr.id]
    pr['family']=[prism_family(t,m) for t,m in zip(pr.target,pr.moa)]; pr['is_family']=pr.family.notna()
    pr['eligible']=(pr.n>=args.min_n)&(pr.n_lineages>=args.min_lineages)
    pr=family_bh(pr)
    pr=pr[['id','drug','screen','dose','target','moa','n','n_lineages','beta','se','t','p','q_global','partial_r2','family','is_family','q_family','eligible']]
    pr.to_csv(args.out_dir/'PRISM_results.csv',index=False)

    gdsc_results={}; gdsc_auc={}; gdsc_model_overlap={}
    sanger_to_i={r.SangerModelID:i for i,r in score.iterrows() if isinstance(r.SangerModelID,str) and r.SangerModelID}
    for platform in ['GDSC1','GDSC2']:
        raw=read_gdsc_xlsx(args.data_dir/f'{platform}_fitted_dose_response_27Oct23.xlsx')
        raw=raw[raw.SANGER_MODEL_ID.isin(sanger_to_i)].copy()
        meta_drug=(raw[['DRUG_ID','DRUG_NAME','PUTATIVE_TARGET','PATHWAY_NAME']].drop_duplicates('DRUG_ID').
                   rename(columns={'DRUG_ID':'drug_id','DRUG_NAME':'drug','PUTATIVE_TARGET':'target','PATHWAY_NAME':'pathway'}))
        model_ids=[s for s in score.SangerModelID if isinstance(s,str) and s and s in set(raw.SANGER_MODEL_ID)]
        model_ids=list(dict.fromkeys(model_ids)); gdsc_model_overlap[platform]=len(model_ids); model_ix=np.array([sanger_to_i[s] for s in model_ids])
        drug_ids=sorted(raw.DRUG_ID.dropna().unique())
        piv=raw.pivot_table(index='SANGER_MODEL_ID',columns='DRUG_ID',values='LN_IC50',aggfunc='median').reindex(index=model_ids,columns=drug_ids)
        Yg=piv.to_numpy(dtype=float); ag=fit_matrix(Yg,lpi[model_ix],pro[model_ix],lineages[model_ix])
        nlin=[len(set(lineages[model_ix][np.isfinite(Yg[:,j])])) for j in range(Yg.shape[1])]
        md=meta_drug.copy(); md['n_lineages']=md.drug_id.map(dict(zip(drug_ids,nlin)))
        d=model_table(ag,drug_ids,md,args.min_n,args.min_lineages)
        d=d[['drug_id','drug','target','pathway','n','n_lineages','beta','se','t','p','q_global','partial_r2','family','is_family','q_family','eligible']]
        d.to_csv(args.out_dir/f'{platform}_results.csv',index=False); gdsc_results[platform]=d
        piva=raw.pivot_table(index='SANGER_MODEL_ID',columns='DRUG_ID',values='AUC',aggfunc='median').reindex(index=model_ids,columns=drug_ids)
        Ya=piva.to_numpy(dtype=float); aa=fit_matrix(Ya,lpi[model_ix],pro[model_ix],lineages[model_ix])
        da=model_table(aa,drug_ids,md,args.min_n,args.min_lineages)
        da=da[['drug_id','drug','target','pathway','n','n_lineages','beta','se','t','p','q_global','partial_r2','family','is_family','q_family','eligible']]
        da.to_csv(args.out_dir/f'{platform}_AUC_results.csv',index=False); gdsc_auc[platform]=da

    pf=pr[pr.is_family & pr.eligible].copy(); pf['norm']=pf.drug.map(normalized_drug)
    cross=[]
    for platform,d in gdsc_results.items():
        df=d[d.is_family & d.eligible].copy(); df['norm']=df.drug.map(normalized_drug)
        for norm,gg in df.groupby('norm'):
            pp=pf[pf.norm.eq(norm)]
            if pp.empty: continue
            for _,prr in pp.iterrows():
                cross.append({'validation_platform':platform,'drug':prr.drug,'family':prr.family,
                              'prism_beta':prr.beta,'prism_p':prr.p,'prism_q_family':prr.q_family,'prism_q_global':prr.q_global,
                              'gdsc_drug':gg.drug.iloc[0],'gdsc_beta_median':float(gg.beta.median()),'gdsc_p_min':float(gg.p.min()),
                              'gdsc_q_family_min':float(gg.q_family.min()),'gdsc_q_global_min':float(gg.q_global.min()),
                              'same_direction':bool(np.sign(prr.beta)==np.sign(gg.beta.median())),'gdsc_n_assays':len(gg)})
    cross_df=pd.DataFrame(cross); cross_df.to_csv(args.out_dir/'cross_platform_family.csv',index=False)

    summary={
      'prism_models':len(prism_ids),'prism_compounds':len(pr),'prism_family_compounds':int(pr.is_family.sum()),
      'prism_global_q05':int((pr.q_global<.05).sum()),'prism_family_q05':int((pr.q_family<.05).sum()),
      'gdsc1_models_overlap':int(gdsc_model_overlap['GDSC1']),'gdsc1_drugs':len(gdsc_results['GDSC1']),
      'gdsc1_global_q05':int((gdsc_results['GDSC1'].q_global<.05).sum()),'gdsc1_family_q05':int((gdsc_results['GDSC1'].q_family<.05).sum()),
      'gdsc2_models_overlap':int(gdsc_model_overlap['GDSC2']),'gdsc2_drugs':len(gdsc_results['GDSC2']),
      'gdsc2_global_q05':int((gdsc_results['GDSC2'].q_global<.05).sum()),'gdsc2_family_q05':int((gdsc_results['GDSC2'].q_family<.05).sum()),
      'exact_cross_platform_rows':len(cross_df)
    }
    a=gdsc_results['GDSC1'][gdsc_results['GDSC1'].drug.str.upper().eq('AS605240')]
    if len(a): summary['AS605240_GDSC1']=a.iloc[0][['beta','p','q_family','q_global','n']].to_dict()
    with open(args.out_dir/'summary.json','w') as f: json.dump(summary,f,indent=2)
    print(json.dumps(summary,indent=2))

if __name__=='__main__': main()
